from pathlib import Path

from docx import Document
from openpyxl import load_workbook


class WorkerListExtractor:
    def extract(self, folder: Path) -> dict:
        excel_files = [path for path in folder.glob("*.xlsx") if not path.name.startswith("~$")]
        if excel_files:
            return self._excel(excel_files[0])
        word_files = list(folder.glob("*.docx"))
        if word_files:
            return self._word(word_files[0])
        return {"status": "missing", "source": None, "file": None, "workers": []}

    def _excel(self, path: Path) -> dict:
        workbook = load_workbook(path, data_only=True, read_only=True)
        return self.from_rows(list(workbook.active.iter_rows(values_only=True)), "excel", path.name)

    def _word(self, path: Path) -> dict:
        rows = [tuple(cell.text for cell in row.cells) for row in Document(path).tables[0].rows]
        return self.from_rows(rows, "word", path.name)

    @staticmethod
    def from_rows(rows: list[tuple], source: str, filename: str) -> dict:
        workers = [
            {"dni": dni, "nombre": f"{first_name} {last_name}", "categoria": category, "horas": int(hours)}
            for dni, first_name, last_name, category, hours in rows[1:]
        ]
        return {"status": "found", "source": source, "file": filename, "workers": workers}
